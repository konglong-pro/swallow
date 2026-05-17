from __future__ import annotations

import json
import wave
import zipfile
from pathlib import Path

import pytest
from typer.testing import CliRunner

from swallow.cli.main import app


FIXTURES = Path(__file__).parent / "fixtures"
pytestmark = pytest.mark.integration

REQUIRED_FIXTURES = [
    "files/simple.txt",
    "files/simple.md",
    "files/simple.docx",
    "files/table.xlsx",
    "files/slides.pptx",
    "files/electronic.pdf",
    "ocr/scanned.pdf",
    "ocr/chinese-scan.pdf",
    "ocr/table-scan.pdf",
    "ocr/screenshot.png",
    "audio/short-zh.wav",
    "audio/short-en.mp3",
    "audio/silent.wav",
    "video/sample.mp4",
    "web/static-article.html",
    "web/dynamic-page.html",
    "web/lazy-load-page.html",
    "web/blocked.html",
    "browser/chatgpt_capture.json",
    "browser/claude_capture.json",
    "browser/gemini_capture.json",
    "archives/chatgpt_export_sample.zip",
    "bad/corrupted.pdf",
    "bad/unsupported.bin",
    "bad/empty.file",
    "bad/path-traversal.zip",
    "bad/zip-slip.zip",
]


def test_fixture_corpus_contains_release_acceptance_files():
    missing = [relative for relative in REQUIRED_FIXTURES if not (FIXTURES / relative).exists()]
    assert missing == []
    for relative in REQUIRED_FIXTURES:
        path = FIXTURES / relative
        if path.name == "empty.file":
            assert path.stat().st_size == 0
        else:
            assert path.stat().st_size > 0, relative
        assert path.stat().st_size < 2_000_000, relative


def test_file_fixtures_are_structurally_readable():
    from docx import Document
    from openpyxl import load_workbook
    from pdfminer.high_level import extract_text
    from pptx import Presentation

    docx = Document(FIXTURES / "files" / "simple.docx")
    assert "Swallow DOCX Fixture" in "\n".join(paragraph.text for paragraph in docx.paragraphs)

    workbook = load_workbook(FIXTURES / "files" / "table.xlsx", read_only=True)
    assert workbook["ingest"]["A1"].value == "input"

    presentation = Presentation(FIXTURES / "files" / "slides.pptx")
    assert len(presentation.slides) == 2

    pdf_text = extract_text(str(FIXTURES / "files" / "electronic.pdf"))
    assert "Swallow electronic PDF fixture" in pdf_text


def test_media_and_archive_fixtures_are_structurally_readable():
    from PIL import Image

    with Image.open(FIXTURES / "ocr" / "screenshot.png") as image:
        assert image.size[0] >= 1000

    with wave.open(str(FIXTURES / "audio" / "silent.wav"), "rb") as audio:
        assert audio.getnchannels() == 1
        assert audio.getframerate() == 16000

    assert (FIXTURES / "audio" / "short-en.mp3").read_bytes()[:3] in {b"ID3", b"\xff\xfb", b"\xff\xf3"}
    assert b"ftyp" in (FIXTURES / "video" / "sample.mp4").read_bytes()[:64]

    with zipfile.ZipFile(FIXTURES / "archives" / "chatgpt_export_sample.zip") as archive:
        payload = json.loads(archive.read("conversations.json").decode("utf-8"))
    assert payload[0]["title"] == "Fixture conversation"

    with zipfile.ZipFile(FIXTURES / "bad" / "zip-slip.zip") as archive:
        names = set(archive.namelist())
    assert "../../evil.txt" in names
    assert "/absolute/path/evil.txt" in names


def test_browser_and_web_fixtures_cover_expected_cases():
    capture = json.loads((FIXTURES / "browser" / "chatgpt_capture.json").read_text(encoding="utf-8"))
    assert capture["platform"] == "chatgpt"
    assert [message["role"] for message in capture["messages"]] == ["user", "assistant"]
    assert "```python" in capture["messages"][1]["content"]

    blocked = (FIXTURES / "web" / "blocked.html").read_text(encoding="utf-8").lower()
    dynamic = (FIXTURES / "web" / "dynamic-page.html").read_text(encoding="utf-8").lower()
    assert "access denied" in blocked
    assert "settimeout" in dynamic


def test_batch_can_run_against_fixed_light_fixture_subset(tmp_path):
    result = CliRunner().invoke(
        app,
        [
            "--store",
            str(tmp_path),
            "batch",
            str(FIXTURES / "files" / "simple.txt"),
            str(FIXTURES / "files" / "simple.md"),
            str(FIXTURES / "bad" / "unsupported.bin"),
            "--json",
        ],
    )

    assert result.exit_code == 2, result.output
    payload = json.loads(result.output)
    summary = json.loads((tmp_path / payload["summary_path"]).read_text(encoding="utf-8"))
    assert summary["total"] == 3
    assert summary["success"] == 2
    assert summary["failed"] == 1
    assert summary["partial"] == 0
