from __future__ import annotations

import json
import math
import os
import shutil
import struct
import subprocess
import wave
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "fixtures"
SAMPLE_TEXT = "Swallow fixture content for ingest regression tests. " * 12


def main() -> int:
    create_dirs()
    create_file_fixtures()
    create_ocr_fixtures()
    create_audio_video_fixtures()
    create_web_fixtures()
    create_browser_fixtures()
    create_archive_fixtures()
    create_bad_fixtures()
    print(f"Generated fixtures under {FIXTURES}")
    return 0


def create_dirs() -> None:
    for relative in ("files", "ocr", "audio", "video", "web", "browser", "archives", "bad"):
        (FIXTURES / relative).mkdir(parents=True, exist_ok=True)


def create_file_fixtures() -> None:
    files = FIXTURES / "files"
    (files / "simple.txt").write_text(SAMPLE_TEXT + "\n", encoding="utf-8")
    (files / "simple.md").write_text("# Swallow Fixture\n\n" + SAMPLE_TEXT + "\n", encoding="utf-8")
    create_docx(files / "simple.docx")
    create_xlsx(files / "table.xlsx")
    create_pptx(files / "slides.pptx")
    create_extractable_pdf(files / "electronic.pdf", "Swallow electronic PDF fixture. " + SAMPLE_TEXT)


def create_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Swallow DOCX Fixture", level=1)
    document.add_paragraph(SAMPLE_TEXT)
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "raw_store"
    table.cell(1, 1).text = "immutable"
    table.cell(2, 0).text = "trace"
    table.cell(2, 1).text = "jsonl"
    document.save(path)


def create_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ingest"
    sheet.append(["input", "worker", "expected"])
    sheet.append(["docx", "markitdown_worker", "markdown"])
    sheet.append(["pdf", "markitdown_worker", "fallback allowed"])
    sheet.append(["image", "paddleocr_worker", "page markdown"])
    workbook.save(path)


def create_pptx(path: Path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Swallow PPTX Fixture"
    slide.placeholders[1].text = "Slide text for MarkItDown extraction."
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Worker Pipeline"
    second.placeholders[1].text = "RawStore -> Router -> WorkerResult -> IngestDocument.md"
    presentation.save(path)


def create_ocr_fixtures() -> None:
    ocr = FIXTURES / "ocr"
    screenshot = ocr / "screenshot.png"
    create_text_image(screenshot, "Swallow OCR screenshot 123", width=1200, height=300)
    create_scanned_pdf_from_text(ocr / "scanned.pdf", "Swallow scanned PDF OCR 123")
    create_scanned_pdf_from_text(ocr / "chinese-scan.pdf", "Swallow 中文 OCR 测试 123", prefer_cjk=True)
    create_table_scan(ocr / "table-scan.pdf")


def create_text_image(path: Path, text: str, *, width: int, height: int, prefer_cjk: bool = False) -> None:
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = load_font(56, prefer_cjk=prefer_cjk)
    draw.text((48, height // 2 - 40), text, fill="black", font=font)
    image.save(path)


def create_scanned_pdf_from_text(path: Path, text: str, *, prefer_cjk: bool = False) -> None:
    image_path = path.with_suffix(".tmp.png")
    create_text_image(image_path, text, width=1200, height=320, prefer_cjk=prefer_cjk)
    try:
        image = Image.open(image_path).convert("RGB")
        image.save(path, "PDF", resolution=200.0)
    finally:
        image_path.unlink(missing_ok=True)


def create_table_scan(path: Path) -> None:
    image = Image.new("RGB", (1200, 520), "white")
    draw = ImageDraw.Draw(image)
    font = load_font(36)
    x0, y0 = 80, 80
    cell_w, cell_h = 340, 95
    rows = [["Field", "Value", "Worker"], ["PDF", "scanned", "PaddleOCR"], ["Table", "text", "Markdown"]]
    for row in range(4):
        y = y0 + row * cell_h
        draw.line((x0, y, x0 + 3 * cell_w, y), fill="black", width=3)
    for col in range(4):
        x = x0 + col * cell_w
        draw.line((x, y0, x, y0 + 3 * cell_h), fill="black", width=3)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            draw.text((x0 + col_index * cell_w + 28, y0 + row_index * cell_h + 28), value, fill="black", font=font)
    image.save(path, "PDF", resolution=200.0)


def create_audio_video_fixtures() -> None:
    audio = FIXTURES / "audio"
    video = FIXTURES / "video"
    short_zh = audio / "short-zh.wav"
    short_en_wav = audio / "short-en.tmp.wav"
    short_en = audio / "short-en.mp3"
    silent = audio / "silent.wav"

    if os.name == "nt" and not synthesize_windows_speech(short_zh, "swallow audio zh fixture one two three"):
        create_tone_wav(short_zh, frequency=440, duration_seconds=1.2)
    if os.name == "nt" and not synthesize_windows_speech(short_en_wav, "swallow audio fixture one two three"):
        create_tone_wav(short_en_wav, frequency=554, duration_seconds=1.2)

    create_silent_wav(silent, duration_seconds=1.0)
    ffmpeg = resolve_ffmpeg()
    if ffmpeg:
        run_ffmpeg([ffmpeg, "-y", "-i", str(short_en_wav), "-codec:a", "libmp3lame", "-q:a", "5", str(short_en)])
        run_ffmpeg(
            [
                ffmpeg,
                "-y",
                "-f",
                "lavfi",
                "-i",
                "color=c=blue:s=320x240:d=1",
                "-f",
                "lavfi",
                "-i",
                "sine=frequency=440:duration=1",
                "-shortest",
                "-pix_fmt",
                "yuv420p",
                str(video / "sample.mp4"),
            ]
        )
    else:
        short_en.write_bytes(short_en_wav.read_bytes())
        (video / "sample.mp4").write_bytes(b"")
    short_en_wav.unlink(missing_ok=True)


def synthesize_windows_speech(path: Path, text: str) -> bool:
    escaped_path = str(path).replace("'", "''")
    escaped_text = text.replace("'", "''")
    command = [
        "powershell",
        "-NoProfile",
        "-Command",
        (
            "$voice = New-Object -ComObject SAPI.SpVoice; "
            "$stream = New-Object -ComObject SAPI.SpFileStream; "
            f"$stream.Open('{escaped_path}', 3); "
            "$voice.AudioOutputStream = $stream; "
            f"$voice.Speak('{escaped_text}') | Out-Null; "
            "$stream.Close()"
        ),
    ]
    completed = subprocess.run(command, capture_output=True, text=True, check=False, timeout=60)
    return completed.returncode == 0 and path.exists() and path.stat().st_size > 0


def create_tone_wav(path: Path, *, frequency: int, duration_seconds: float) -> None:
    sample_rate = 16000
    frames = int(sample_rate * duration_seconds)
    with wave.open(str(path), "wb") as output:
        output.setnchannels(1)
        output.setsampwidth(2)
        output.setframerate(sample_rate)
        for index in range(frames):
            value = int(12000 * math.sin(2 * math.pi * frequency * index / sample_rate))
            output.writeframes(struct.pack("<h", value))


def create_silent_wav(path: Path, *, duration_seconds: float) -> None:
    sample_rate = 16000
    frames = int(sample_rate * duration_seconds)
    with wave.open(str(path), "wb") as output:
        output.setnchannels(1)
        output.setsampwidth(2)
        output.setframerate(sample_rate)
        output.writeframes(b"\x00\x00" * frames)


def resolve_ffmpeg() -> str | None:
    found = shutil.which("ffmpeg")
    if found:
        return found
    try:
        import imageio_ffmpeg

        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        return None


def run_ffmpeg(command: list[str]) -> None:
    completed = subprocess.run(command, capture_output=True, text=True, check=False, timeout=60)
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr.strip() or completed.stdout.strip() or "ffmpeg failed")


def create_web_fixtures() -> None:
    web = FIXTURES / "web"
    (web / "static-article.html").write_text(
        """
<!doctype html>
<html><head><title>Static Article Fixture</title></head>
<body>
<article>
<h1>Static Article Fixture</h1>
<p>Swallow static article content for URL and HTML conversion tests.</p>
<p>RawStore remains immutable and trace events explain every worker step.</p>
</article>
</body></html>
""".strip()
        + "\n",
        encoding="utf-8",
    )
    (web / "dynamic-page.html").write_text(
        """
<!doctype html>
<html><head><title>Dynamic Page Fixture</title></head>
<body>
<main id="app">Loading...</main>
<script>
setTimeout(() => {
  document.querySelector("#app").innerHTML = "<h1>Dynamic Page Fixture</h1><p>Rendered content after JavaScript.</p>";
}, 10);
</script>
</body></html>
""".strip()
        + "\n",
        encoding="utf-8",
    )
    (web / "lazy-load-page.html").write_text(
        """
<!doctype html>
<html><head><title>Lazy Load Fixture</title></head>
<body>
<h1>Lazy Load Fixture</h1>
<div style="height:1200px">Scroll down</div>
<section id="lazy">Lazy loaded content fixture for Playwright scrolling.</section>
</body></html>
""".strip()
        + "\n",
        encoding="utf-8",
    )
    (web / "blocked.html").write_text(
        """
<!doctype html>
<html><head><title>Blocked Fixture</title></head>
<body><h1>Access denied</h1><p>Status 403. Please log in to continue.</p></body></html>
""".strip()
        + "\n",
        encoding="utf-8",
    )


def create_browser_fixtures() -> None:
    browser = FIXTURES / "browser"
    for platform in ("chatgpt", "claude", "gemini"):
        payload = {
            "platform": platform,
            "url": f"https://{platform}.example.com/c/fixture",
            "title": f"{platform.title()} capture fixture",
            "captured_at": "2026-05-16T12:00:00+08:00",
            "messages": [
                {"role": "user", "content": "Create an ingest fixture with a link: https://example.com"},
                {"role": "assistant", "content": "```python\nprint('swallow')\n```\nFixture response preserved verbatim."},
            ],
            "raw_dom": f"<html><body><main>{platform} fixture</main></body></html>",
        }
        (browser / f"{platform}_capture.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def create_archive_fixtures() -> None:
    archive_path = FIXTURES / "archives" / "chatgpt_export_sample.zip"
    conversations = [
        {
            "title": "Fixture conversation",
            "create_time": 1.0,
            "update_time": 3.0,
            "mapping": {
                "root": {"message": None},
                "user": {
                    "message": {
                        "author": {"role": "user"},
                        "create_time": 1.5,
                        "content": {"content_type": "text", "parts": ["Please ingest this export fixture."]},
                    }
                },
                "assistant": {
                    "message": {
                        "author": {"role": "assistant"},
                        "create_time": 2.0,
                        "content": {"content_type": "text", "parts": ["Export fixture parsed through WorkerResult."]},
                    }
                },
            },
        }
    ]
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("conversations.json", json.dumps(conversations, ensure_ascii=False, indent=2))


def create_bad_fixtures() -> None:
    bad = FIXTURES / "bad"
    (bad / "corrupted.pdf").write_bytes(b"%PDF-1.4\nnot a valid pdf body\n%%EOF\n")
    (bad / "unsupported.bin").write_bytes(b"\x00\x01unsupported swallow fixture\x02\x03")
    (bad / "empty.file").write_bytes(b"")
    with zipfile.ZipFile(bad / "path-traversal.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("../../evil.txt", "path traversal fixture")
        archive.writestr("nested/../../../evil.txt", "nested traversal fixture")
    with zipfile.ZipFile(bad / "zip-slip.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("../../evil.txt", "zip slip fixture")
        archive.writestr("/absolute/path/evil.txt", "absolute path fixture")


def create_extractable_pdf(path: Path, text: str) -> None:
    lines = wrap_text(text, width=72)[:18]
    commands = ["BT", "/F1 12 Tf", "72 740 Td"]
    for index, line in enumerate(lines):
        if index:
            commands.append("0 -18 Td")
        commands.append(f"({escape_pdf_text(line)}) Tj")
    commands.append("ET")
    stream = "\n".join(commands).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    write_pdf_objects(path, objects)


def wrap_text(text: str, *, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if len(candidate) > width and current:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def escape_pdf_text(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def write_pdf_objects(path: Path, objects: list[bytes]) -> None:
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{index} 0 obj\n".encode("ascii"))
        content.extend(obj)
        content.extend(b"\nendobj\n")

    xref_offset = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    content.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    content.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    path.write_bytes(content)


def load_font(size: int, *, prefer_cjk: bool = False):
    candidates = []
    if prefer_cjk:
        candidates.extend(
            [
                Path("C:/Windows/Fonts/msyh.ttc"),
                Path("C:/Windows/Fonts/simsun.ttc"),
                Path("C:/Windows/Fonts/simhei.ttf"),
            ]
        )
    candidates.extend(
        [
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("C:/Windows/Fonts/segoeui.ttf"),
            Path("C:/Windows/Fonts/calibri.ttf"),
        ]
    )
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


if __name__ == "__main__":
    raise SystemExit(main())
